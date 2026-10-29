"""Bounded, non-executing CSV/XLSX/ODS row readers (no office runtime)."""
from __future__ import annotations
import csv
from pathlib import Path
import zipfile
from lxml import etree as ET
from openpyxl import load_workbook
from .package import MAX_ARCHIVE, PackageError, safe_entries

T = '{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
X = '{urn:oasis:names:tc:opendocument:xmlns:text:1.0}'
O = '{urn:oasis:names:tc:opendocument:xmlns:office:1.0}'
LINK = '{http://www.w3.org/1999/xlink}href'
MAX_DATA_ROWS = 500_000
MAX_COLUMNS = 256


def _count(element, name, maximum):
    value = element.get(T + name, '1')
    if not value.isdecimal() or not 1 <= int(value) <= maximum:
        raise PackageError('ODS repetition resource limit: ' + name)
    return int(value)


def _text(element):
    if element.tag == X + 's':
        value = element.get(X + 'c', '1')
        if not value.isdecimal() or not 1 <= int(value) <= 4096:
            raise PackageError('ODS text space resource limit')
        return ' ' * int(value)
    if element.tag == X + 'tab': return '\t'
    if element.tag == X + 'line-break': return '\n'
    return (element.text or '') + ''.join(_text(c) + (c.tail or '') for c in element)


def ods_rows(path: Path):
    """Yield (sheet, original 1-based row, trimmed string cells).

    Read every sheet, even if a caller only needs CM. Never expand a million
    blank rows or 16,384 empty trailing cells produced by LibreOffice.
    """
    from .adapters import xml_root
    if path.stat().st_size > MAX_ARCHIVE: raise PackageError('ODS archive size limit')
    with zipfile.ZipFile(path) as archive:
        entries = safe_entries(archive, package=False)
        if archive.read('mimetype') != b'application/vnd.oasis.opendocument.spreadsheet':
            raise PackageError('Not an ODS spreadsheet')
        for entry in entries:
            if entry.filename != 'content.xml' and entry.filename.endswith('.xml'):
                root = xml_root(archive.read(entry))
                if any(ET.QName(e).localname == 'encryption-data' for e in root.iter()):
                    raise PackageError('Encrypted ODS prohibited')
        with archive.open('content.xml') as source:
            parser = ET.iterparse(source, events=('start', 'end'), resolve_entities=False,
                                  load_dtd=False, no_network=True, huge_tree=False)
            sheet = None; row_number = 0; emitted = 0; seen = set(); depth = 0
            for event, element in parser:
                if event == 'start':
                    depth += 1
                    if depth > 64: raise PackageError('ODS XML depth limit')
                    if element.getroottree().docinfo.doctype:
                        raise PackageError('DTD/entities prohibited')
                    if element.tag == T + 'table':
                        if sheet is not None: raise PackageError('Nested ODS table prohibited')
                        sheet = element.get(T + 'name'); row_number = 0
                        if not sheet or sheet in seen: raise PackageError('Missing/duplicate ODS sheet')
                        seen.add(sheet)
                    if T + 'formula' in element.attrib:
                        raise PackageError('Spreadsheet formulas prohibited')
                    if element.get(LINK, '').strip() and not element.get(LINK).startswith('#'):
                        raise PackageError('External ODS links prohibited')
                    if ET.QName(element).localname in {'scripts', 'script', 'object', 'object-ole', 'event-listeners'}:
                        raise PackageError('Active ODS content prohibited')
                    continue
                if element.tag == T + 'table-row':
                    repeat = _count(element, 'number-rows-repeated', 1_048_576)
                    cells = []; pending = 0; column_count = 0
                    for cell in element:
                        if cell.tag not in {T + 'table-cell', T + 'covered-table-cell'}:
                            raise PackageError('Unknown ODS row child')
                        n = _count(cell, 'number-columns-repeated', 16_384)
                        column_count += n
                        if column_count > 16_384: raise PackageError('ODS column resource limit')
                        value = '\n'.join(_text(p) for p in cell.findall(X + 'p'))
                        if not value and cell.get(O + 'value-type'):
                            value = cell.get(O + 'string-value', cell.get(O + 'value', cell.get(O + 'date-value', cell.get(O + 'boolean-value', ''))))
                        if value:
                            if len(cells) + pending + n > MAX_COLUMNS:
                                raise PackageError('ODS nonempty column resource limit')
                            cells.extend([''] * pending); cells.extend([value] * n); pending = 0
                        else: pending += n
                    if sheet is None: raise PackageError('ODS row outside sheet')
                    if row_number + repeat > 1_048_576: raise PackageError('ODS row resource limit')
                    if cells:
                        emitted += repeat
                        if emitted > MAX_DATA_ROWS: raise PackageError('ODS nonempty row resource limit')
                        for number in range(row_number + 1, row_number + repeat + 1):
                            yield sheet, number, list(cells)
                    row_number += repeat
                    element.clear()
                    while element.getprevious() is not None: del element.getparent()[0]
                elif element.tag == T + 'table':
                    sheet = None
                    element.clear()
                depth -= 1
            if parser.root is None or parser.root.tag != O + 'document-content':
                raise PackageError('Unknown ODS content root')


def spreadsheet_rows(path: Path, sheet: str | None = None):
    if path.stat().st_size > MAX_ARCHIVE: raise PackageError('Translation file size limit')
    suffix = path.suffix.lower()
    if suffix == '.ods':
        if not sheet: raise PackageError('ODS requires explicit sheet; CM and PCS must not be mixed')
        found = False
        for name, line, row in ods_rows(path):
            if name == sheet:
                found = True; yield line, row
        if not found: raise PackageError('ODS sheet not found: ' + sheet)
    elif suffix == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as source:
            for line, row in enumerate(csv.reader(source), 1):
                if line > MAX_DATA_ROWS: raise PackageError('CSV row resource limit')
                if len(row) > MAX_COLUMNS: raise PackageError('CSV column resource limit')
                yield line, row
    elif suffix == '.xlsx':
        from .adapters import xml_root
        with zipfile.ZipFile(path) as archive:
            for entry in safe_entries(archive, package=False):
                if entry.filename.lower().endswith(('.xml', '.rels')): xml_root(archive.read(entry))
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            selected = workbook[sheet or workbook.sheetnames[0]]
            if selected.max_row and selected.max_row > MAX_DATA_ROWS: raise PackageError('XLSX row resource limit')
            if selected.max_column and selected.max_column > MAX_COLUMNS: raise PackageError('XLSX column resource limit')
            for number, cells in enumerate(selected, 1):
                if any(c.data_type == 'f' for c in cells): raise PackageError('Spreadsheet formulas prohibited')
                yield number, ['' if c.value is None else str(c.value) for c in cells]
        finally: workbook.close()
    else: raise PackageError('Unsupported translation format; use UTF-8 CSV, XLSX or ODS')


def mapped_records(path: Path, columns: dict):
    required = {'code', 'title_zh_tw'}
    if not required <= columns.keys(): raise PackageError('Explicit code/title_zh_tw column mapping required')
    rows = spreadsheet_rows(path, columns.get('sheet'))
    try:
        _, header = next(rows)
        if not header or len(set(header)) != len(header): raise PackageError('Empty/duplicate headers')
        if not all(columns[k] in header for k in required): raise PackageError('Mapped column missing; update mapping JSON')
        for number, row in rows:
            if not any(row): continue
            if len(row) > len(header) and any(row[len(header):]): raise PackageError('Unexpected extra source columns')
            padded = (row + [''] * len(header))[:len(header)]
            yield number, dict(zip(header, padded))
    except StopIteration as exc: raise PackageError('Empty translation table') from exc
    finally: rows.close()
